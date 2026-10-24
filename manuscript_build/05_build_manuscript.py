#!/usr/bin/env python3
"""Build the JEI manuscript .docx on the official JEI Author Manuscript Template (styles, header/footer,
continuous line numbering, 1-inch margins kept), then render a PDF with LibreOffice."""
import re, sys, copy, subprocess
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent; OUT = HERE.parent
sys.path.insert(0, str(HERE)); import manuscript_text as T
CIT = re.compile(r"\[@([a-z0-9_,]+)\]"); ORDER = []
def _scan(t):
    for m in CIT.finditer(t):
        for k in m.group(1).split(","):
            assert k in T.REFS, k
            if k not in ORDER: ORDER.append(k)
for it in T.INTRO + T.RESULTS + T.DISCUSSION + T.METHODS:
    if isinstance(it, str): _scan(it)
def _fmt(nums):
    nums = sorted(nums); out = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ", ".join(str(n) for n in nums[i:j + 1])); i = j + 1
    return "(" + ", ".join(out) + ")"
def cite(t): return CIT.sub(lambda m: _fmt([ORDER.index(k) + 1 for k in m.group(1).split(",")]), t)
T.INTRO = [cite(x) if isinstance(x, str) else x for x in T.INTRO]; T.RESULTS = [cite(x) if isinstance(x, str) else x for x in T.RESULTS]
T.DISCUSSION = [cite(x) if isinstance(x, str) else x for x in T.DISCUSSION]; T.METHODS = [cite(x) if isinstance(x, str) else x for x in T.METHODS]
REFLIST = [T.REFS[k] for k in ORDER]
print("refs used:", len(REFLIST), "unused:", [k for k in T.REFS if k not in ORDER])
TEMPLATE = OUT / "references_check" / "JEI_Author_Manuscript_Template_downloaded.docx"
DOCX = OUT / "manuscript" / "Mamidala_JEI_manuscript_v5_simulated_missingness_CKD.docx"
APPX = OUT / "supplement" / "Mamidala_JEI_v4_Appendix.docx"
RES = OUT / "repo" / "results"; FIG = OUT / "repo" / "figures"

doc = Document(str(TEMPLATE))
body = doc.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial") if st.element.rPr is not None and st.element.rPr.rFonts is not None else None

TOK = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\^\{.+?\}|_\{.+?\})")
def add_runs(p, text, bold=False, italic=False, size=11):
    for part in TOK.split(text):
        if not part: continue
        b, i, sup, sub = bold, italic, False, False
        if part.startswith("**"): part, b = part[2:-2], True
        elif part.startswith("*"): part, i = part[1:-1], True
        elif part.startswith("^{"): part, sup = part[2:-1], True
        elif part.startswith("_{"): part, sub = part[2:-1], True
        r = p.add_run(part); r.bold = b; r.italic = i
        r.font.name = "Arial"; r.font.size = Pt(size); r.font.color.rgb = RGBColor(0, 0, 0)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial"); r._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        if sup: r.font.superscript = True
        if sub: r.font.subscript = True
    return p

def para(text="", bold=False, italic=False, align=None, size=11, spacing=1.5, indent=False, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.line_spacing = spacing; pf.space_before = Pt(0); pf.space_after = Pt(0)
    if indent: pf.first_line_indent = Inches(0.3)
    if align: p.alignment = align
    if keep: pf.keep_with_next = True
    add_runs(p, text, bold, italic, size); return p

def page_break():
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)

def heading(t): return para(t, bold=True, keep=True)
def subheading(t): return para(t, italic=True, keep=True)

def section(items):
    for it in items:
        if isinstance(it, tuple): subheading(it[1])
        else: para(it, indent=True)

# ---- Title page ----
para(T.TITLE, bold=True)
para(""); para(T.AUTHORS, bold=True)
for a in T.AFFIL: para(a)
para(""); para("Student Authors", bold=True)
for s_ in T.STUDENTS: para(s_)
para(""); para("**KEYWORDS:** " + T.KEYWORDS)
para(""); para("**OVERVIEW:** " + T.OVERVIEW)
page_break()
# ---- Summary ----
heading("SUMMARY"); para(T.SUMMARY)
page_break()
heading("INTRODUCTION"); section(T.INTRO)
heading("RESULTS"); section(T.RESULTS)
heading("DISCUSSION"); section(T.DISCUSSION)
heading("MATERIALS AND METHODS"); section(T.METHODS)
heading("ACKNOWLEDGMENTS"); para(T.ACK, indent=True)
heading("REFERENCES")
for i, r in enumerate(REFLIST, 1):
    para(f"{i}. " + r)   # numbered, no hanging indent (JEI)

# ---- Figures ----
page_break(); heading("Figures and Figure Titles/Captions")
for k, (fname, title, cap) in enumerate(T.FIGS):
    if k: page_break()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / f"{fname}.png"), width=Inches(min(6.5, __import__("PIL.Image").Image.open(str(FIG / f"{fname}.png")).size[0] / 600)))
    q = para("", spacing=1.5); add_runs(q, title, bold=True); add_runs(q, cap)

# ---- Tables ----
def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.find(qn("w:tcBorders"))
    if b is None: b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, val in kw.items():
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), val.get("val", "single")); e.set(qn("w:sz"), str(val.get("sz", 4)))
        e.set(qn("w:color"), "000000"); b.append(e)

def make_table(df, size, widths, bold_rows=(), group_rows=()):
    t = doc.add_table(rows=df.shape[0] + 1, cols=df.shape[1]); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, c in enumerate(df.columns):
        cell = t.cell(0, j); cell.text = ""; pp = cell.paragraphs[0]; add_runs(pp, str(c), bold=True, size=size)
        pp.paragraph_format.line_spacing = 1.0
        set_cell_border(cell, top={"sz": 12}, bottom={"sz": 8})
    for i in range(df.shape[0]):
        for j in range(df.shape[1]):
            cell = t.cell(i + 1, j); cell.text = ""; pp = cell.paragraphs[0]
            v = str(df.iat[i, j]).replace("-0.", "−0.").replace("-", "−") if j > 0 else str(df.iat[i, j])
            add_runs(pp, v, bold=(i in group_rows), italic=False, size=size); pp.paragraph_format.line_spacing = 1.0
            pp.paragraph_format.space_before = Pt(1); pp.paragraph_format.space_after = Pt(1)
            if j > 0: pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == df.shape[0] - 1: set_cell_border(cell, bottom={"sz": 12})
    # fixed layout with explicit grid so Word and LibreOffice honour the column widths
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None: tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:w"), str(int(sum(widths) * 1440))); tw.set(qn("w:type"), "dxa")
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))): gc.set(qn("w:w"), str(int(widths[j] * 1440)))
    for row in t.rows:
        for j, w in enumerate(widths): row.cells[j].width = Inches(w)
    for i in group_rows:
        row = t.rows[i + 1]; m = row.cells[0].merge(row.cells[-1])
        for extra in m.paragraphs[1:]: extra._element.getparent().remove(extra._element)
    return t

page_break(); heading("Tables with Titles/Captions")
t1 = pd.read_csv(RES / "table1_cohort.csv").fillna("")
t1.columns = ["Characteristic", "Overall", "Low income\n(PIR < 1.3)", "Middle income\n(1.3 ≤ PIR < 3.5)", "High income\n(PIR ≥ 3.5)"]
t1["Characteristic"] = t1["Characteristic"].str.replace("kg/m²", "kg/m^{2}").str.replace("1.73 m²", "1.73 m^{2}")
t1["Characteristic"] = t1["Characteristic"].str.replace("^  ", "    ", regex=True)
make_table(t1, 9.5, [2.3, 1.05, 1.05, 1.1, 1.0])
q = para("", spacing=1.5); add_runs(q, T.TABLE1_TITLE, bold=True); add_runs(q, T.TABLE1_CAP)

page_break()
t2 = pd.read_csv(RES / "table2_primary.csv").fillna("")
groups = [k for k in range(len(t2)) if t2.iloc[k, 1] == ""]
rows = t2.values.tolist()
make_table(pd.DataFrame(rows, columns=t2.columns), 9, [2.1, 1.1, 1.1, 1.1, 1.1], group_rows=tuple(groups))
q = para("", spacing=1.5); add_runs(q, T.TABLE2_TITLE, bold=True); add_runs(q, T.TABLE2_CAP)

page_break()
t3 = pd.read_csv(RES / "table3_sensitivity.csv")
rows, groups = [], []
for a, sub in t3.groupby("Analysis", sort=False):
    groups.append(len(rows)); rows.append([a, "", "", ""])
    for r in sub.itertuples(index=False): rows.append(["    " + r[1], r[2], r[3], r[4]])
make_table(pd.DataFrame(rows, columns=["Analysis / repair", "Low income", "High income", "Low minus high"]), 8.5, [2.6, 1.3, 1.3, 1.3], group_rows=tuple(groups))
q = para("", spacing=1.5); add_runs(q, T.TABLE3_TITLE, bold=True); add_runs(q, T.TABLE3_CAP)

DOCX.parent.mkdir(exist_ok=True); doc.save(str(DOCX)); print("saved", DOCX)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(DOCX.parent), str(DOCX)],
               check=True, capture_output=True, timeout=240)
print("pdf ok")

